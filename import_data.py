"""
Script import data historis v2 — hapus data lama lalu import ulang.
Letakkan file ini di root project (sejajar dengan manage.py), lalu jalankan:
    python import_data_v2.py

Ganti nama file Excel di variabel FILE_EXCEL di bawah.
"""
import os
import django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

import openpyxl
from rekrutmen_guru.models import (
    BidangStudi, PeriodeRekrutmen, Kriteria,
    Kandidat, NilaiKriteria, HasilSeleksi
)

# ── Ganti nama file Excel kamu di sini ────────────────────────────────────────
FILE_EXCEL = 'Data_Rekrutmen_Training_Fixed.xlsx'
NAMA_SHEET = 'Data Training'

# ── Mapping nama kriteria di Excel → nama di database ─────────────────────────
KOLOM_KE_KRITERIA = {
    'kualifikasi_akademik' : 'Kualifikasi Akademik',
    'iq'                   : 'IQ',
    'tes_psikotes'         : 'Tes Psikotes',
    'nilai_akademik'       : 'Nilai Akademik',
    'pengalaman_tahun'     : 'Pengalaman Mengajar',
    'micro_teaching'       : 'Micro Teaching',
    'komitmen'             : 'Komitmen (Wawancara)',
    'ruhiyah'              : 'Ruhiyah (Keagamaan)',
}

def hapus_data_historis():
    """Hapus semua data historis (kandidat yang punya label)."""
    historis = Kandidat.objects.filter(label__isnull=False)
    jumlah = historis.count()
    if jumlah == 0:
        print("Tidak ada data historis yang perlu dihapus.")
        return

    konfirmasi = input(f"⚠ Akan menghapus {jumlah} kandidat historis beserta nilai dan hasil seleksinya. Lanjutkan? (y/n): ")
    if konfirmasi.lower() != 'y':
        print("Dibatalkan.")
        exit()

    historis.delete()
    print(f"✅ {jumlah} kandidat historis berhasil dihapus.\n")


def import_data():
    wb = openpyxl.load_workbook(FILE_EXCEL, data_only=True)

    if NAMA_SHEET not in wb.sheetnames:
        print(f"❌ Sheet '{NAMA_SHEET}' tidak ditemukan di file {FILE_EXCEL}")
        print(f"   Sheet yang tersedia: {wb.sheetnames}")
        exit()

    ws = wb[NAMA_SHEET]

    # Ambil semua kriteria dari database
    kriteria_db = {k.nama: k for k in Kriteria.objects.all()}

    # Buat periode default untuk data historis
    periode, _ = PeriodeRekrutmen.objects.get_or_create(nama='Data Historis')

    sukses, skip, error = 0, 0, []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if len(row) < 13:
            continue

        nama, jenjang, jurusan, ipk, kual_ak, pengalaman, iq, \
        tes_psikotes, nilai_akademik, komitmen, ruhiyah, micro, label = row[:13]

        if not nama:
            continue

        # Filter label
        if label == 'Diterima':
            label_db = 'diterima'
        elif label == 'Tidak Diterima':
            label_db = 'tidak_diterima'
        else:
            skip += 1
            continue

        try:
            # Buat atau update bidang studi
            bidang, _ = BidangStudi.objects.get_or_create(
                nama=str(jurusan).strip() if jurusan else 'Umum'
            )

            # Update atau buat kandidat
            kandidat, created = Kandidat.objects.update_or_create(
                nama=nama,
                defaults={
                    'bidang_studi'       : bidang,
                    'jenjang_pendidikan' : str(jenjang).strip() if jenjang else 'S1',
                    'ipk'                : float(ipk) if ipk else None,
                    'pengalaman_mengajar': int(pengalaman) if pengalaman else 0,
                    'periode'            : periode,
                    'label'              : label_db,
                    'status_tahap_1'     : 'hadir',
                    'status_tahap_2'     : 'hadir',
                    'status_tahap_3'     : 'hadir',
                    'status_tahap_4'     : 'hadir',
                }
            )

            # Update nilai per kriteria
            nilai_map = {
                'kualifikasi_akademik' : float(kual_ak)       if kual_ak       is not None else None,
                'iq'                   : float(iq)             if iq            is not None else None,
                'tes_psikotes'         : float(tes_psikotes)   if tes_psikotes  is not None else None,
                'nilai_akademik'       : float(nilai_akademik) if nilai_akademik is not None else None,
                'pengalaman_tahun'     : float(pengalaman)     if pengalaman    is not None else 0,
                'micro_teaching'       : float(micro)          if micro         is not None else None,
                'komitmen'             : float(komitmen)       if komitmen      is not None else None,
                'ruhiyah'              : float(ruhiyah)        if ruhiyah       is not None else None,
            }

            for kolom, nilai in nilai_map.items():
                if nilai is None:
                    continue
                nama_kriteria = KOLOM_KE_KRITERIA[kolom]
                if nama_kriteria not in kriteria_db:
                    continue
                NilaiKriteria.objects.update_or_create(
                    kandidat=kandidat,
                    kriteria=kriteria_db[nama_kriteria],
                    defaults={'nilai': nilai}
                )

            aksi = "Dibuat" if created else "Diupdate"
            sukses += 1

        except Exception as e:
            error.append(f"Baris {i} ({nama}): {e}")

    print(f"✅ Berhasil import  : {sukses} kandidat")
    print(f"⏭  Dilewati         : {skip} kandidat (label tidak jelas)")
    if error:
        print(f"\n❌ Error ({len(error)}):")
        for e in error:
            print(f"   {e}")
    else:
        print(f"❌ Error            : 0")
    print(f"\n💡 Jangan lupa jalankan training ulang model di halaman Training Model!")


if __name__ == '__main__':
    print(f"📂 File: {FILE_EXCEL}")
    print(f"📋 Sheet: {NAMA_SHEET}\n")
    hapus_data_historis()
    import_data()